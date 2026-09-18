from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os
from datetime import date

REPORTS_DIR = "reports"
os.makedirs(REPORTS_DIR, exist_ok=True)

SIGNAL = RGBColor(0x10, 0xB9, 0x81)
DARK = RGBColor(0x18, 0x18, 0x1B)
LIGHT = RGBColor(0xF4, 0xF4, 0xF5)
MUTED = RGBColor(0x71, 0x71, 0x7A)


def _bullets_slide(prs, title, items):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = items[0]
    for item in items[1:]:
        p = body.add_paragraph()
        p.text = item
    return slide


def generate_ppt_report(dataset_name: str, summary: str, kpis: list, recommendations: list,
                        anomalies: list = [], sample_data: list = []) -> str:
    filename = os.path.join(REPORTS_DIR, f"report_{dataset_name}.pptx")
    prs = Presentation()

    # Slide 1 — Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"BI Report — {dataset_name}"
    subtitle = slide.placeholders[1].text_frame
    subtitle.text = "Autonomous Business Intelligence Platform"
    p = subtitle.add_paragraph()
    p.text = date.today().isoformat()

    # Slide 2 — Executive Summary
    _bullets_slide(prs, "Executive Summary", [summary])

    # Slide 3 — Key Metrics
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Key Metrics"
    rows = [["Measure", "Total", "Average", "Min", "Max"]]
    for k in kpis:
        rows.append([k["measure"], k["total"], k["average"], k["minimum"], k["maximum"]])
    table_shape = slide.shapes.add_table(
        rows=len(rows), cols=5, left=Inches(0.6), top=Inches(1.8),
        width=Inches(11.0), height=Inches(0.4 * len(rows)),
    )
    tbl = table_shape.table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.text_frame.paragraphs[0].font.size = Pt(12)
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = SIGNAL
                cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.text_frame.paragraphs[0].font.bold = True
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                cell.text_frame.paragraphs[0].font.color.rgb = DARK

    # Slide 4 — Anomalies
    anomaly_items = ([f"Anomalous value {a.get('value', '?')} — z-score {a.get('z_score', '?')}" for a in anomalies]
                     or ["No anomalies detected."])
    _bullets_slide(prs, "Anomalies", anomaly_items)

    # Slide 5 — Recommendations
    rec_items = ([f"{i}. {rec}" for i, rec in enumerate(recommendations, start=1)] or ["No recommendations available."])
    _bullets_slide(prs, "Recommendations", rec_items)

    # Slide 6 — Appendix (first 10 rows)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Appendix — Sample Data"
    if sample_data:
        headers = list(sample_data[0].keys())
        rows = [headers]
        for row in sample_data[:10]:
            rows.append([row.get(h, "") for h in headers])
        table_shape = slide.shapes.add_table(
            rows=len(rows), cols=len(headers), left=Inches(0.6), top=Inches(1.8),
            width=Inches(11.0), height=Inches(0.3 * len(rows)),
        )
        tbl = table_shape.table
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = str(val)
                cell.text_frame.paragraphs[0].font.size = Pt(9)
                if r == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = SIGNAL
                    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    cell.text_frame.paragraphs[0].font.bold = True
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = LIGHT if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                    cell.text_frame.paragraphs[0].font.color.rgb = DARK
    else:
        slide.placeholders[1].text_frame.text = "No sample data available."

    prs.save(filename)
    return filename